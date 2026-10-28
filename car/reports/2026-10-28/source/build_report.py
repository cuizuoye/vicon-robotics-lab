from pathlib import Path
import hashlib
import json
import math
from urllib.parse import quote

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[3]
WORK = ROOT.parent / '.local/report-rebuild'
WORK.mkdir(parents=True, exist_ok=True)
OUT = WORK / 'Vicon标定杆遥操作ROS全向小车项目技术说明.docx'
FONT = 'Noto Sans CJK SC'
REG = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
BOLD = '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'


def diagram():
    im = Image.new('RGB', (1650, 960), 'white')
    d = ImageDraw.Draw(im)
    navy, blue, line, ink = '#15334B', '#EAF2F8', '#3A718E', '#172D3B'

    def text(x, y, s, size=27, bold=False, fill=ink, anchor='mm'):
        d.text((x, y), s, font=ImageFont.truetype(BOLD if bold else REG, size),
               fill=fill, anchor=anchor)

    def box(x, y, w, h, title, lines, fill=blue):
        d.rounded_rectangle((x, y, x+w, y+h), radius=15, fill=fill,
                            outline=line, width=3)
        text(x+w/2, y+38, title, 32, True)
        for i, s in enumerate(lines):
            text(x+w/2, y+84+i*36, s, 26)

    def arrow(points, color=line, width=5):
        d.line(points, fill=color, width=width, joint='curve')
        (a, b), (x, y) = points[-2:]
        angle = math.atan2(y-b, x-a)
        tri = [(x, y), (x-17*math.cos(angle-.45), y-17*math.sin(angle-.45)),
               (x-17*math.cos(angle+.45), y-17*math.sin(angle+.45))]
        d.polygon(tri, fill=color)

    box(70, 40, 440, 150, '空间主操作器 Sign', ['T形五点杆', '水平位移与相对旋转'])
    box(780, 40, 750, 150, 'Vicon 统一空间测量',
        ['同时跟踪 Sign 与 car01', '位置  姿态  时间戳  跟踪有效性'], '#DCEBF4')
    arrow([(510, 115), (780, 115)])
    text(645, 83, '光学跟踪', 25)
    arrow([(1070, 190), (1070, 280)])
    d.rectangle((790, 215, 1350, 253), fill='white')
    text(1070, 234, 'Sign目标输入与car01实际位姿反馈', 25)

    d.rounded_rectangle((70, 280, 1570, 560), radius=16, fill='#F5F8FA',
                        outline='#9AAEBB', width=3)
    text(110, 315, 'Linux PC 与 ROS 2    位置外环约50 Hz', 30, True, anchor='lm')
    box(105, 360, 410, 160, '相对位姿映射', ['起点与朝向锚定', '1比1 XY 与相对yaw'])
    box(620, 360, 400, 160, '低延迟连续伺服', ['速度前馈与位置纠偏', '限速  限加速度'])
    box(1125, 360, 410, 160, '监督器与末级安全门', ['来源与新鲜度校验', '故障停车  禁止自动恢复'])
    arrow([(515, 440), (620, 440)])
    arrow([(1020, 440), (1125, 440)])

    arrow([(1330, 520), (1330, 615), (280, 615), (280, 690)])
    d.rectangle((570, 586, 1120, 642), fill='white')
    text(845, 614, 'UDP MC-Link 指令与反向遥测', 26)
    box(70, 690, 420, 185, 'Tenda 与 ESP32 S3', ['专用Wi-Fi通信链路', '报文校验与事件收发', '不执行轨迹控制'])
    box(610, 690, 420, 185, 'STM32F407', ['全向运动学与分档约束', '编码器速度内环100 Hz', '四轮独立前馈与PI'])
    box(1150, 690, 420, 185, '驱动板与四轮小车', ['AT8236与MG310电机', '平移  侧移  原地旋转', '编码器返回实际轮速'])
    arrow([(490, 760), (610, 760)])
    text(550, 724, 'UART', 25)
    arrow([(1030, 752), (1150, 752)])
    text(1090, 716, 'PWM', 24)
    arrow([(1150, 833), (1030, 833)], '#647987', 4)
    text(1090, 872, '编码器', 23)
    arrow([(1570, 780), (1613, 780), (1613, 115), (1530, 115)], '#177C80', 5)
    text(1360, 657, '实际整车运动反馈', 25, fill='#176C70')
    text(825, 929, 'Vicon闭合整车位置外环    编码器闭合四轮速度内环', 29, True)
    path = WORK / 'system_architecture.png'
    im.save(path)
    return path


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.67), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.76), Inches(.76)
sec.footer_distance = Inches(.25)
for grid in sec._sectPr.findall(qn('w:docGrid')):
    sec._sectPr.remove(grid)


def set_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:'+key), FONT)
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:'+key), None)
    for borders in style.element.findall('.//'+qn('w:pBdr')):
        borders.getparent().remove(borders)
    pp = style.element.get_or_add_pPr()
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    pp.append(snap)


set_font(doc.styles['Normal'], 11)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = Pt(16)
normal.space_after = Pt(7)
normal.widow_control = True
for name, size in [('Title', 24), ('Subtitle', 11), ('Heading 1', 17),
                   ('Heading 2', 12.5), ('Heading 3', 11.5), ('Caption', 10)]:
    set_font(doc.styles[name], size, name not in ('Subtitle', 'Caption'))
    pf = doc.styles[name].paragraph_format
    pf.space_before = Pt(11 if name.startswith('Heading') else 0)
    pf.space_after = Pt(7)
    pf.keep_with_next = name not in ('Caption',)
    pf.line_spacing = Pt(31 if name == 'Title' else (23 if name == 'Heading 1' else 18))
set_font(doc.styles['Footer'], 9)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Vicon 标定杆遥操作项目  |  ')
f = OxmlElement('w:fldSimple')
f.set(qn('w:instr'), 'PAGE')
footer._p.append(f)

doc.core_properties.title = '基于Vicon动捕与标定杆遥操作的ROS全向小车技术说明'
doc.core_properties.subject = '项目汇报与分级验收状态'
doc.core_properties.author = 'Vicon ROS小车项目'
doc.core_properties.keywords = 'Vicon ROS2 遥操作 麦克纳姆轮 位置闭环 验收'


def p(text='', bold=False, size=None):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return para


def h(text, level=1):
    return doc.add_heading(text, level)


def bullet(text):
    para = p('•  '+text)
    para.paragraph_format.left_indent = Inches(.08)
    para.paragraph_format.space_after = Pt(5)
    return para


def page(title):
    doc.add_page_break()
    h(title)


def table(headers, rows, widths, size=10.3):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Inches(width)
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = text
    tblpr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:'+edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '5')
        el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    tblpr.append(borders)
    for values in rows:
        cells = t.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width)
            cell.text = str(value)
    for ridx, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        trpr.append(no_split)
        if ridx == 0:
            repeat = OxmlElement('w:tblHeader')
            trpr.append(repeat)
        for cidx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, amount in [('top', 100), ('bottom', 100), ('left', 115), ('right', 115)]:
                el = OxmlElement('w:'+side)
                el.set(qn('w:w'), str(amount))
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '183B52' if ridx == 0 else ('F1F5F8' if ridx % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.line_spacing = Pt(14.5)
                if cidx == 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(size)
                    if ridx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    p().paragraph_format.space_after = Pt(0)
    return t


def math_run(text):
    r = OxmlElement('m:r')
    props = OxmlElement('m:rPr')
    sty = OxmlElement('m:sty')
    sty.set(qn('m:val'), 'p')
    props.append(sty)
    r.append(props)
    rr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Cambria Math')
    fonts.set(qn('w:hAnsi'), 'Cambria Math')
    rr.append(fonts)
    r.append(rr)
    el = OxmlElement('m:t')
    el.text = text
    r.append(el)
    return r


def sub(base, index):
    el = OxmlElement('m:sSub')
    e = OxmlElement('m:e')
    e.append(math_run(base))
    s = OxmlElement('m:sub')
    s.append(math_run(index))
    el.extend([e, s])
    return el


def equation(parts):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(7)
    para.paragraph_format.line_spacing = Pt(18)
    mp = OxmlElement('m:oMathPara')
    eq = OxmlElement('m:oMath')
    for part in parts:
        eq.append(math_run(part) if isinstance(part, str) else part)
    mp.append(eq)
    para._p.append(mp)
    return para


def link(para, text, target):
    rid = para.part.relate_to(target, RT.HYPERLINK, is_external=True)
    a = OxmlElement('w:hyperlink')
    a.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    col = OxmlElement('w:color')
    col.set(qn('w:val'), '205C80')
    pr.append(col)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '19')
    pr.append(sz)
    r.append(pr)
    txt = OxmlElement('w:t')
    txt.text = text
    r.append(txt)
    a.append(r)
    para._p.append(a)


# Page 1
doc.add_paragraph('基于Vicon动捕与标定杆遥操作的\nROS全向小车技术说明', 'Title')
doc.add_paragraph('项目汇报  2026年10月28日  版本1.0', 'Subtitle')
h('项目结论')
p('本项目已建立由Vicon统一空间测量、ROS 2整车位姿闭环和STM32四轮速度闭环组成的移动机器人系统。我们已完成实车标定杆平移跟随，并在独立的早期方案中完成矩形循环及受扰回轨。平移与旋转叠加、高加速度同步控制已有仿真验证；远程浏览器控制和长时间连续稳定性仍需单独验收。')
h('项目目标与系统边界', 2)
p('操作者手持T形五点标定杆Sign，通过自然的水平移动和转动控制麦克纳姆轮小车car01。当前采用1比1相对位移映射，并具备1比1相对航向映射的控制链。目标是直观展示动捕定位支持的人机空间交互、全向跟随和外部扰动纠偏，而不是仅演示遥控电机转动。')
h('Vicon在系统中的核心作用')
h('为人和车建立共同空间参考', 2)
p('Vicon在同一坐标系中测量Sign和car01。控制器无需依靠操作者估计距离，即可将杆的相对动作转换为车的目标位置与朝向。系统同时保留源时间戳和跟踪有效性，用于判断输入能否继续参与控制。')
h('观察并纠正整车实际运动', 2)
p('编码器反映轮子转动，Vicon反映车体实际位姿。小车打滑、侧移不足或被人为推离时，位置外环能观察实际偏差并进行纠正。该作用是矩形受扰恢复和空间遥操作的共同基础。')
h('提供可视化与验收依据', 2)
p('目标轨迹、车体实测轨迹和标定杆输入可以在同一时间轴下比较，用于分析跟随误差、换向拖后、遮挡停车和恢复流程。定位精度及曝光到执行器的完整延迟尚未经过专门计量，不能直接引用厂商能力代替本系统实测结果。')
p('当前可报告的成果是受限条件下的闭环演示能力。取消固定围栏不等于具备避障能力；配置速度上限也不等于已测得持续运行速度。', bold=True)

# Page 2
page('整体架构与技术分工')
p('图1展示控制指令向下执行、轮速遥测反向返回，以及Vicon重新测量整车运动的完整闭环。操作台负责交互，不以界面绘图坐标代替真实定位。')
pp = doc.add_paragraph()
pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
pp.paragraph_format.line_spacing = 1.0
pp.paragraph_format.keep_with_next = True
run = pp.add_run()
run.add_picture(str(diagram()), width=Inches(6.96))
inline = run._r.xpath('.//wp:docPr')[0]
inline.set('descr', 'Sign经Vicon向ROS 2提供目标，car01经Vicon提供整车反馈。ROS 2映射、伺服和安全门通过ESP32与STM32驱动四轮小车，编码器闭合轮速内环。')
doc.add_paragraph('图1  Vicon位置外环与编码器速度内环的分层架构', 'Caption')
table(['层级', '技术与职责'], [
    ['感知与ROS', 'Vicon接收节点、PoseStamped与PointStamped、rclpy及DDS。处理刚体位姿、标记点几何、时间戳和来源。'],
    ['控制与安全', '相对位姿映射、连续伺服、轨迹恢复、监督状态机及独立末级输出门。电脑控制周期配置为50 Hz。'],
    ['通信与执行', 'Tenda专用Wi-Fi、ESP32-S3、UDP MC-Link、CRC32及会话序号。UART采用921600 bit/s与COBS分帧。'],
    ['底层与硬件', 'STM32F407 C固件、定时轮控100 Hz、霍尔编码器、AT8236双路驱动板及MG310减速电机。'],
    ['交互与工程', '本机Tkinter操作台与RViz；Python回归、真实ROS节点测试、C故障注入、ARM构建及固件读回校验。浏览器远程端独立开发。'],
], [1.13, 5.83])
p('系统硬件采用四个MG310霍尔编码器电机、两块AT8236驱动板及3S电池。已提供的车体外廓为190 mm × 210 mm，整车质量约1.4 kg。ESP32只承担通信转发，不运行轨迹跟踪或轮速控制。', size=10.3)

# Page 3
page('关键算法与空间遥操作')
h('相对平移映射与原点时序', 2)
p('设杆和车在本次实际使能时的位置分别为杆零点和车零点。控制器先将杆的水平位移变换到启动车头坐标，再按比例映射回世界坐标；1比1映射下等价于下式。准备阶段不捕获运动原点，允许操作者先拿起、调整杆，再由健康的首次使能帧建立原点。')
equation([sub('p', 'target'), ' = ', sub('p', 'car0'), ' + s (', sub('p', 'Sign'), ' − ', sub('p', 'Sign0'), ')    s = 1'])
p('式中各p只包含水平面XY分量。杆回到杆零点，目标就回到车零点；杆的高度Z不驱动车辆平移。启用边界限幅的历史版本在边缘不再严格保持线性比例，不能与后续无固定围栏入口混为一谈。')
h('基于T形几何的相对旋转', 2)
p('横杆使用Sign1、Sign5、Sign2，长柄使用Sign3、Sign4。系统用五点相对几何核对点号和刚体姿态，定义车头方向为朝横杆外侧、背离长柄；将该方向投影到水平面获取航向。左右方向已用真实杆约正反90度只读动作核对。[E03]')
equation([sub('ψ', 'target'), ' = ', sub('ψ', 'car0'), ' + Δ', sub('ψ', 'Sign')])
p('相对转角连续处理跨越正负180度的情况，平移与旋转可同时输入。旋转仿真以横杆中心Sign5作为跟随参考点。抬杆或翻腕不直接作为绕竖直轴的旋转指令；当水平朝向不可可靠定义、几何不一致或定位失效时，应停止接受该输入。')
h('低延迟连续伺服', 2)
p('控制器不再堆叠位置低通。它用最近约40 ms源时间戳样本估计目标速度，以速度前馈跟随手部趋势，再用位置误差反馈消除漂移。下式为控制思想的简化表达，实际实现还包含参考点引导、车体坐标变换、运动状态增益切换及输出约束。')
equation([sub('v', 'cmd'), ' = limit (', sub('v', 'ff'), ' + ', sub('K', 'p'), ' (', sub('p', 'ref'), ' − ', sub('p', 'car'), '))'])
p('前馈不外推或伪造新位姿。缺少新样本时先撤去速度前馈，定位达到超时条件后由安全层停车。电脑连续伺服与底层加速斜坡经过协调，避免多层慢整形叠加，但正常加速度限制与换向中性保护仍保留。')
h('参数匹配的仿真结论', 2)
p('在同一峰值0.50 m/s平顺往返模型中，原0.20 m/s与0.08 m/s²限制下位置误差RMS为25.98 cm；0.70 m/s与2.00 m/s²组合为4.94 cm，换向延迟中位数约由1.43 s降至0.16 s。这些结果来自指定二维模型，不是实车误差或相机到车轮延迟。[E04]')

# Page 4
page('全向执行与分层安全')
h('麦克纳姆轮逆运动学与轮速闭环', 2)
p('小车按ROS习惯定义前为x正向、左为y正向、逆时针为正偏航。对X型麦克纳姆轮，车体速度可以同时包含前进速度、侧移速度和偏航角速度。四个轮子的目标角速度按下式分配，其中L为轮系半轴距与半轮距之和，r为轮半径。')
equation([sub('ω', 'FL'), ' = (', sub('v', 'x'), ' − ', sub('v', 'y'), ' − L', sub('ω', 'z'), ') / r'])
equation([sub('ω', 'FR'), ' = (', sub('v', 'x'), ' + ', sub('v', 'y'), ' + L', sub('ω', 'z'), ') / r'])
equation([sub('ω', 'RL'), ' = (', sub('v', 'x'), ' + ', sub('v', 'y'), ' − L', sub('ω', 'z'), ') / r'])
equation([sub('ω', 'RR'), ' = (', sub('v', 'x'), ' − ', sub('v', 'y'), ' + L', sub('ω', 'z'), ') / r'])
p('轮速超限时进行共同缩放，尽量保留合成运动方向。STM32用每轮编码器和前馈加PI闭环控制PWM，并保留积分抗饱和。正常反转与过零按轮独立处理，避免单轮换向被错误理解为整车异常；真实堵转、异常反馈及超时仍可触发停车。[E11]')
h('矩形轨迹与外部扰动恢复', 2)
p('矩形控制器对当前有向边做投影，计算前视目标、沿边进度和横向误差。受扰时保持当前边，减小沿边推进并加强横向纠偏，回到容差后继续运行。它不会简单选择全局最近边，因此可避免被推近矩形对边时错误跳段。历史方案允许的恢复偏差有上限，不承诺任意位置均能恢复。[E02]')
h('监督器与独立末级门', 2)
bullet('准入时核对定位、通信握手、固件身份、无故障与零输出；检查通过并不等于自动运动，仍需新的操作员启动动作。')
bullet('运行中校验源身份、序号、时间戳、指令限值和执行器反馈。监督器产生指令后，末级门再次独立决定是否允许输出。')
bullet('遮挡、过期、异常姿态、通信中断、来源冲突或硬故障触发停车；数据回来不自动续跑，必须重新建立有效会话。')
p('本机候选的代表性时效门限为定位200 ms、指令100 ms、操作员租约150 ms和MCU反馈300 ms；浏览器只读候选另保留100 ms回应保护。它们是不同层的失效判据，不能相加成端到端延迟，也不能等同于机械完全停稳时间。[E08–E10]')

# Page 5
page('已完成成果与验证分级')
p('成果按具体能力和具体版本分级。真实杆驱动软件车属于仿真验证，编译与回归通过属于离线软件证据；两者都不替代真实车体执行验收。')
h('实车及真实输入已验证', 2)
table(['能力', '已观察结果', '证据与范围'], [
    ['实车平移跟随', '1比1全向跟随、前移与返回、停车；操作者确认同步性较理想。', 'E01。保存基线的运行约19.52 s，以松键结束。0.70 m/s和0.80 m/s²为配置上限；不是完整3×3 m遍历或实测速率认证。'],
    ['历史矩形与受扰回轨', '完成整圈、圈间重装填及一次随机干扰后的回轨继续运行。', 'E02。对应独立历史方案；记录对后续4WGB四角版本仍保留待验证项，不能泛化到全部后续固件。'],
    ['真实杆方向核对', '确认横杆点号与正方向；正反约90度旋转均可读出。', 'E03。只读真实输入验证，不是电机旋转或精密角度计量验收。'],
    ['较长实车跟随', '一次记录前60 s未出现MCU反馈超时或门控故障，故障位为0。', 'E08。延长后约98.6 s发生ROS指令保护停车；不能据此宣称120 s或长期连续运行已通过。'],
], [1.15, 2.32, 3.49], size=10.1)
h('仿真与离线软件已通过', 2)
table(['验证对象', '结论', '不覆盖的内容'], [
    ['平移同步与恢复', '模型匹配对照完成；真实Sign驱动软件车的同步、换向、松键、遮挡及手动恢复通过。[E04、E13]', '0.70 m/s与2.00 m/s²为软件车参数，不证明本车带载加速度、电流、温升及停车距离。'],
    ['真实杆驱动旋转仿真', 'A2平移与旋转跟随、换向及操作员停止通过；用户反馈同步性较好。[E05]', 'A2记录未验收遮挡与手动恢复；不能借用早期平移仿真的恢复结果。'],
    ['ROS参考帧修复', '历史保留、精确匹配与队列相关修复完成真实ROS软件用例验证。[E09]', '候选未完成物理运动验收，不能视为本机默认入口已整体升级。'],
    ['浏览器低积压只读端', '本机Chrome严格60 s只读测试通过，控制与画面分流、画面有界积压。[E10]', 'Windows远程待验收，ROS软件跟随未合格；该入口只允许只读。'],
], [1.15, 2.95, 2.86], size=10.1)

# Page 6
page('当前边界与后续验收计划')
h('版本与参数必须成套说明', 2)
table(['方案', '主要参数', '当前资格'], [
    ['实车平移基线', '1比1；0.70 m/s；0.80 m/s²；固定3×3 m目标区；20 s入口。', '仅已观察短程动作通过。不是完整区域与长期稳定性验收。[E01]'],
    ['后续本机演示入口', '1比1；120 s会话；可取消固定围栏；平移上限仍为0.70与0.80；旋转另有限速。', '入口和软件能力存在；不将高加速度A2仿真参数自动继承为实车能力。[E06–E08、E12]'],
    ['A2旋转仿真', '平移0.70 m/s与2.00 m/s²；偏航2.00 rad/s与12.00 rad/s²。', '真实Sign驱动软件车的正常跟随与操作员停止通过。不是实车许可。[E05]'],
], [1.35, 3.13, 2.48], size=10.1)
h('仍待完成的关键项', 2)
table(['优先项', '验证方法', '建议完成判据'], [
    ['长时间连续稳定性', '隔离测试ROS参考帧修复候选，记录来源匹配、队列、指令与MCU反馈年龄。', '完成整段120 s及多次启停；无非预期超时或来源丢失，并核对停车和重新准备。'],
    ['实车旋转与组合运动', '独立验收左右转、跨零角、平移加旋转、回正及停车。', '方向与角度正确，无异常冲击；用同时间位姿统计误差，不用仿真角速度代替实测。'],
    ['高加速度及电气负载', '分档带载试验，测停车包络、压降、电流和温升。', '取得实际峰值与持续负载证据；普通万用表只能初筛，不能认证瞬态峰值和热安全。'],
    ['完整Vicon覆盖演示', '测绘有效覆盖、遮挡与边缘质量，建立空旷运动路线。', '在代表位置重复验证定位与停车；覆盖有效不等于周围无障碍。'],
    ['Windows远程操作', '先完成浏览器严格只读，再做隔离ROS软件跟随，最后另行实车验收。', '保留100 ms保护。只读、仿真与真机分别通过，不通过时不开放运动入口。'],
], [1.35, 2.85, 2.76], size=10.1)
p('必要边界：当前没有障碍物检测或力反馈；完整覆盖区未测绘；最坏制动距离、峰值电流、温升和真实端到端延迟未获认证。取消围栏后，现场空旷条件与可立即断电仍是运行前提。', bold=True, size=10.5)

# Page 7
page('证据索引与汇报口径')
p('以下索引定位到项目保存记录。文件链接指向公开仓库中的脱敏记录；文档转移到其他电脑后，可直接打开链接或按编号检索。记录与原始日志应随正式验收资料保留。', size=10.2)

evidence = [
    ('E01', '实车平移跟随接受基线', 'docs/wand_servo_4ws8_2026-10-26/demo_3x3_20s/accepted_20261026/ACCEPTED_BASELINE.json'),
    ('E02', '历史连续矩形与随机干扰回轨记录', 'docs/CONTINUOUS_3X2_VICON_RECOVERY_2026-10-21.md'),
    ('E03', 'T形杆点号与方向只读核对', 'docs/wand_rotation_preview_2026-10-26/DIRECTION_ACCEPTANCE.json'),
    ('E04', '纯仿真速度加速度匹配对照报告', 'docs/wand_servo_matching_sim_2026-10-23/REPORT.md'),
    ('E05', '真实Sign驱动A2软件车接受记录', 'docs/wand_rotation_accel_sim_2026-10-26/A2_FOLLOW_ACCEPTANCE.json'),
    ('E06', '平移旋转实机控制链离线准备', 'docs/wand_rotation_hardware_2026-10-26/prepared/PREPARATION_RESULT.json'),
    ('E07', '无固定围栏入口的软件准备结果', 'docs/wand_tracking_workspace_2026-10-27/prepared/PREPARATION_RESULT.json'),
    ('E08', '实车60秒观察及98秒后保护停车记录', 'docs/wand_feedback_300ms_startup_2026-10-27/GROUND_60S_RESULT.json'),
    ('E09', 'ROS参考帧保留匹配与队列修复候选', 'docs/wand_ros_bound_history_2026-10-27/prepared/PREPARATION_RESULT.json'),
    ('E10', '浏览器快速只读候选验收结果', 'docs/wand_browser_fast_2026-10-28/prepared/PREPARATION_RESULT.json'),
    ('E11', '4WS8分档固件设计与资格边界', 'stm32/mecanum_f407/candidates/servo_4ws8/README.md'),
    ('E12', '本机自动自检与单击启动准备结果', 'docs/wand_tracking_autostart_2026-10-27/prepared/PREPARATION_RESULT.json'),
    ('E13', '真实Sign驱动平移软件车与遮挡恢复验收', 'docs/wand_servo_matching_sim_2026-10-23/LIVE_ACCEPTANCE.md'),
]
manifest = []
for eid, label, rel in evidence:
    src = ROOT / rel
    assert src.is_file(), src
    row = p(eid+'  '+label, bold=True, size=10.3)
    row.paragraph_format.keep_with_next = True
    row.paragraph_format.space_after = Pt(1)
    ref = doc.add_paragraph()
    ref.paragraph_format.space_after = Pt(6)
    link(ref, src.name+'  查看源记录', 'https://github.com/cuizuoye/vicon-robotics-lab/blob/main/car/'+quote(rel, safe='/'))
    manifest.append(dict(id=eid, title=label, path=str(src), sha256=hashlib.sha256(src.read_bytes()).hexdigest()))

h('项目汇报建议表述', 2)
p('本项目利用Vicon同时测量人手标定杆与移动小车，在统一坐标系内实现相对动作映射，并通过ROS 2位置外环和编码器速度内环驱动车辆。实车已展示平移跟随和历史轨迹受扰恢复；平移叠加旋转及更高加速度同步控制已有仿真结果。后续将通过独立实机验收补齐长期稳定性、旋转性能、停车包络及电气负载证据。', size=10.5)
p('术语口径：RMS表示同时间样本的均方根误差；配置上限不是实测峰值；控制周期不是端到端延迟；软件零输出不是机械完全停稳；4WS8是本项目固件版本标签，并非四轮转向机构。', size=9.8)

doc.save(OUT)
(WORK / 'evidence_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
print(OUT)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables), 'evidence', len(manifest))
